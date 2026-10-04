"""Master-resume text extraction (spec §13).

Deterministic text extraction from PDF (PyMuPDF) and DOCX (python-docx).
No LLM here — structuring happens in agents/profile_agent.py.
"""

from __future__ import annotations

from pathlib import Path

from config.logging_config import get_logger

logger = get_logger(__name__)

# Minimum characters for the extracted text to be considered usable.
MIN_RESUME_CHARS = 100


class ResumeParsingError(Exception):
    """Raised when a resume file cannot be read or yields too little text."""


def extract_text_from_pdf(path: str | Path) -> str:
    """Extract plain text from a PDF using PyMuPDF."""
    import pymupdf  # PyMuPDF

    text_parts: list[str] = []
    with pymupdf.open(str(path)) as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts).strip()


def extract_text_from_docx(path: str | Path) -> str:
    """Extract plain text from a DOCX using python-docx (paragraphs + tables)."""
    from docx import Document

    doc = Document(str(path))
    lines: list[str] = [p.text for p in doc.paragraphs]
    # Include table cell text, which resumes often use for layout.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    lines.append(cell.text)
    return "\n".join(line for line in lines if line is not None).strip()


def extract_text(path: str | Path) -> str:
    """Dispatch extraction by file extension and validate content length."""
    path = Path(path)
    ext = path.suffix.lower()
    logger.info("Extracting text from resume (ext=%s)", ext)

    try:
        if ext == ".pdf":
            text = extract_text_from_pdf(path)
        elif ext == ".docx":
            text = extract_text_from_docx(path)
        else:
            raise ResumeParsingError(f"Unsupported resume type: {ext}")
    except ResumeParsingError:
        raise
    except Exception as exc:  # noqa: BLE001 - surface a clean error to the UI
        logger.error("Failed to extract resume text: %s", exc)
        raise ResumeParsingError(f"Could not read resume: {exc}") from exc

    if len(text) < MIN_RESUME_CHARS:
        raise ResumeParsingError(
            "Extracted resume text is too short to parse reliably "
            f"({len(text)} chars). The file may be scanned/image-based."
        )

    logger.info("Extracted %d characters of resume text", len(text))
    return text
