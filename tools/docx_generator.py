"""DOCX rendering (python-docx).

Produces an editable DOCX that mirrors the PDF template: strictly black-and-white,
single-column, standard headings. Driven by config/formatting.yaml.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from config.logging_config import get_logger
from models.resume_models import CustomizedResume

logger = get_logger(__name__)

# ATS-safe sans-serif that pairs with Helvetica in the PDF.
DOCX_FONT = "Arial"
BLACK = RGBColor(0, 0, 0)


def _set_margins(section, margins: dict) -> None:
    section.top_margin = Inches(margins.get("top", 0.5))
    section.bottom_margin = Inches(margins.get("bottom", 0.5))
    section.left_margin = Inches(margins.get("left", 0.6))
    section.right_margin = Inches(margins.get("right", 0.6))


def _add_bottom_border(paragraph) -> None:
    """Add a thin black bottom border to a paragraph (section-heading rule)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")        # ~0.75pt
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    p_pr.append(borders)


def _run(paragraph, text: str, *, size: int, bold: bool = False, italic: bool = False):
    run = paragraph.add_run(text)
    run.font.name = DOCX_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    return run


def _heading(doc, text: str, size: int) -> None:
    p = doc.add_paragraph()
    p.space_before = Pt(8)
    _run(p, text.upper(), size=size, bold=True)
    _add_bottom_border(p)


def _bullet(doc, text: str, size: int) -> None:
    p = doc.add_paragraph(style="List Bullet")
    _run(p, text, size=size)


def render_docx(resume: CustomizedResume, out_path: str | Path, fmt: dict) -> Path:
    """Render the resume to an editable, black-and-white DOCX."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    font = fmt.get("font", {})
    body = font.get("body_size", 10)
    heading = font.get("heading_size", 12)
    name_size = font.get("name_size", 16)

    doc = Document()
    # Base style.
    normal = doc.styles["Normal"]
    normal.font.name = DOCX_FONT
    normal.font.size = Pt(body)
    normal.font.color.rgb = BLACK
    _set_margins(doc.sections[0], fmt.get("margins", {}))

    p = resume.personal

    # Header (centered name + contact).
    if p.name:
        hp = doc.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(hp, p.name, size=name_size, bold=True)
    contact = [p.email, p.phone, p.location, p.linkedin, p.github, p.portfolio]
    contact = [c for c in contact if c]
    if contact:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(cp, "  |  ".join(contact), size=body - 1)

    order = fmt.get("section_order", [
        "summary", "skills", "experience", "projects", "education", "certifications",
    ])
    for extra in ("achievements", "languages", "hobbies"):
        if extra not in order:
            order.append(extra)

    for section in order:
        if section == "summary" and resume.summary:
            _heading(doc, "Summary", heading)
            sp = doc.add_paragraph()
            sp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _run(sp, resume.summary, size=body)

        elif section == "skills" and resume.skills:
            _heading(doc, "Skills", heading)
            _run(doc.add_paragraph(), ", ".join(resume.skills), size=body)

        elif section == "experience" and resume.experience:
            _heading(doc, "Experience", heading)
            for e in resume.experience:
                title = ", ".join(x for x in [e.title, e.company] if x)
                _run(doc.add_paragraph(), title, size=body, bold=True)
                meta = " | ".join(
                    x for x in [e.location,
                                " - ".join(d for d in [e.start_date, e.end_date] if d)]
                    if x
                )
                if meta:
                    _run(doc.add_paragraph(), meta, size=body - 1, italic=True)
                for b in e.bullets:
                    _bullet(doc, b, body)

        elif section == "projects" and resume.projects:
            _heading(doc, "Projects", heading)
            for pr in resume.projects:
                _run(doc.add_paragraph(), pr.name, size=body, bold=True)
                sub = []
                if pr.technologies:
                    sub.append("Technologies: " + ", ".join(pr.technologies))
                if pr.link:
                    sub.append(pr.link)
                if sub:
                    _run(doc.add_paragraph(), " | ".join(sub), size=body - 1, italic=True)
                if pr.description:
                    _run(doc.add_paragraph(), pr.description, size=body)
                for b in pr.bullets:
                    _bullet(doc, b, body)

        elif section == "education" and resume.education:
            _heading(doc, "Education", heading)
            for ed in resume.education:
                line = ", ".join(x for x in [ed.degree, ed.institution] if x)
                _run(doc.add_paragraph(), line, size=body, bold=True)
                years = " - ".join(d for d in [ed.start_year, ed.end_year] if d)
                meta = " | ".join(
                    x for x in [years, (f"Grade: {ed.grade}" if ed.grade else "")] if x
                )
                if meta:
                    _run(doc.add_paragraph(), meta, size=body - 1, italic=True)

        elif section == "certifications" and resume.certifications:
            _heading(doc, "Certifications", heading)
            for c in resume.certifications:
                _bullet(doc, c, body)

        elif section == "achievements" and resume.achievements:
            _heading(doc, "Achievements", heading)
            for a in resume.achievements:
                _bullet(doc, a, body)

        elif section == "languages" and resume.languages:
            _heading(doc, "Languages", heading)
            _run(doc.add_paragraph(), ", ".join(resume.languages), size=body)

        elif section == "hobbies" and resume.hobbies:
            _heading(doc, "Hobbies", heading)
            _run(doc.add_paragraph(), ", ".join(resume.hobbies), size=body)

    doc.save(str(out_path))
    logger.info("DOCX rendered: %s", out_path.name)
    return out_path
