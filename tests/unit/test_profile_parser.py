"""Unit tests for the Master Profile Agent and resume text extraction.

Uses a deterministic fake LLM — no live model calls (spec §34).
"""

from __future__ import annotations

import pytest

from models.resume_models import MasterProfile


class _FakeStructured:
    """Stands in for llm.with_structured_output(MasterProfile)."""

    def __init__(self, result):
        self._result = result

    def invoke(self, messages):
        # Confirm the agent passes system + human messages.
        assert any(role == "system" for role, _ in messages)
        assert any(role == "human" for role, _ in messages)
        return self._result


class _FakeLLM:
    def __init__(self, result):
        self._result = result

    def with_structured_output(self, schema):
        assert schema is MasterProfile
        return _FakeStructured(self._result)


def test_structure_resume_text_with_model_result():
    from agents.profile_agent import structure_resume_text

    expected = MasterProfile(
        personal={"name": "Jane Doe", "email": "jane@example.com"},
        summary="Data engineer with 5 years of experience.",
        skills=["Python", "Spark"],
    )
    profile = structure_resume_text("raw resume text", llm=_FakeLLM(expected))
    assert profile.personal.name == "Jane Doe"
    assert profile.skills == ["Python", "Spark"]


def test_structure_resume_text_normalizes_dict_result():
    """Agent accepts a dict (some provider paths) and validates it."""
    from agents.profile_agent import structure_resume_text

    dict_result = {
        "personal": {"name": "John"},
        "summary": "",
        "skills": ["SQL"],
    }
    profile = structure_resume_text("raw", llm=_FakeLLM(dict_result))
    assert isinstance(profile, MasterProfile)
    assert profile.personal.name == "John"
    assert profile.skills == ["SQL"]


def test_extract_text_rejects_short_content(tmp_path):
    from tools.resume_parser import ResumeParsingError, extract_text

    f = tmp_path / "tiny.pdf"
    f.write_bytes(b"too short")
    with pytest.raises(ResumeParsingError):
        extract_text(f)


def test_extract_text_rejects_unknown_extension(tmp_path):
    from tools.resume_parser import ResumeParsingError, extract_text

    f = tmp_path / "resume.txt"
    f.write_text("x" * 200)
    with pytest.raises(ResumeParsingError):
        extract_text(f)


def test_normalize_location():
    from agents.profile_agent import normalize_location

    assert normalize_location("Hyderabad,") == "Hyderabad"
    assert normalize_location("Hyderabad ,India") == "Hyderabad, India"
    assert normalize_location("  Bangalore  ") == "Bangalore"
    assert normalize_location("") == ""


def test_completeness_flags_missing_and_grade_in_grade_field():
    from agents.profile_agent import check_profile_completeness

    incomplete = MasterProfile(
        personal={"name": "Jane"},  # missing email/phone/location
        skills=["Python"],
        experience=[{"title": "Engineer", "company": "Acme"}],  # no dates
        education=[
            {"degree": "B.Tech IT", "institution": "VNR VJIET"}  # no end_year/grade
        ],
    )
    missing = check_profile_completeness(incomplete)
    assert any("Email" in m for m in missing)
    assert any("dates for experience" in m for m in missing)
    assert any("Grade/CGPA" in m for m in missing)
    assert any("Completion year" in m for m in missing)

    complete = MasterProfile(
        personal={
            "name": "Jane",
            "email": "j@x.com",
            "phone": "123",
            "location": "Hyderabad",
        },
        skills=["Python"],
        experience=[
            {"title": "Eng", "company": "Acme", "start_date": "2021", "end_date": "Present"}
        ],
        education=[
            {"degree": "B.Tech", "institution": "VNR", "end_year": "2020", "grade": "9.11"}
        ],
    )
    assert check_profile_completeness(complete) == []


def test_education_grade_is_separate_from_year():
    """A CGPA now lives in `grade`, not a year field (bug fix)."""
    from models.resume_models import EducationItem

    edu = EducationItem(
        degree="B.Tech IT", institution="VNR VJIET", end_year="2020", grade="9.11"
    )
    assert edu.grade == "9.11"
    assert edu.end_year == "2020"
    # The old `year` field no longer exists.
    assert not hasattr(edu, "year")


def test_extract_text_from_docx_roundtrip(tmp_path):
    """Create a real DOCX and extract its text deterministically."""
    from docx import Document

    from tools.resume_parser import extract_text

    doc = Document()
    doc.add_paragraph("Jane Doe — Data Engineer")
    doc.add_paragraph("Skills: Python, SQL, Spark, Airflow, dbt, Snowflake")
    doc.add_paragraph("Experience: Built ETL pipelines at Acme Corp 2020-2024.")
    path = tmp_path / "resume.docx"
    doc.save(path)

    text = extract_text(path)
    assert "Jane Doe" in text
    assert "Spark" in text
